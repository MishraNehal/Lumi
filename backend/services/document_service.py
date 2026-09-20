import os
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutureTimeoutError
from langchain_community.document_loaders import (
    PyPDFLoader, TextLoader,
    UnstructuredWordDocumentLoader, UnstructuredPowerPointLoader, UnstructuredExcelLoader,
)
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from backend.database.vectorstore import get_vector_store
from backend.services.ocr_service import ingest_ocr_document


# ── Fallback glue-code parsers (used only if the LangChain/Unstructured loader fails) ──

def _fallback_load_docx(file_path: str) -> list:
    from docx import Document as DocxDocument
    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    text = "\n\n".join(paragraphs)
    return [Document(page_content=text, metadata={})] if text.strip() else []


def _fallback_load_pptx(file_path: str) -> list:
    from pptx import Presentation
    prs = Presentation(file_path)
    documents = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        texts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        texts.append(row_text)
        slide_text = "\n".join(texts)
        if slide_text.strip():
            documents.append(Document(page_content=slide_text, metadata={"slide": i + 1}))
    return documents


def _fallback_load_excel(file_path: str) -> list:
    import pandas as pd
    sheets = pd.read_excel(file_path, sheet_name=None)
    documents = []
    for sheet_name, df in sheets.items():
        text = df.to_string(index=False)
        if text.strip():
            documents.append(Document(page_content=text, metadata={"sheet": sheet_name}))
    return documents


# ── Loader dispatch: try LangChain/Unstructured first (with a hard timeout, since
#    it can HANG rather than raise — e.g. libmagic/NLTK issues on Windows — and a
#    plain try/except can't catch a hang), fall back to glue code either way. ──

_LOADER_TIMEOUT_SECONDS = 20


def _run_with_timeout(fn, timeout=_LOADER_TIMEOUT_SECONDS):
    """
    Runs fn() in a background thread; raises TimeoutError if it hangs past `timeout`s.
    Deliberately does NOT use ThreadPoolExecutor as a context manager — exiting a
    `with` block waits for the worker thread to finish, which would block on a hang
    and defeat the entire point of the timeout.
    """
    executor = ThreadPoolExecutor(max_workers=1)
    future = executor.submit(fn)
    try:
        return future.result(timeout=timeout)
    except FutureTimeoutError:
        executor.shutdown(wait=False)  # abandon the hung thread, don't wait for it
        raise TimeoutError(f"Loader did not finish within {timeout}s (likely hung).")
    finally:
        if not future.cancelled() and future.done():
            executor.shutdown(wait=False)


def _load_docx(file_path: str) -> list:
    try:
        documents = _run_with_timeout(lambda: UnstructuredWordDocumentLoader(file_path).load())
        if any(d.page_content.strip() for d in documents):
            return documents
        raise ValueError("Unstructured loader returned empty content.")
    except Exception as e:
        print(f"⚠️ UnstructuredWordDocumentLoader failed/timed out, using fallback parser: {e}")
        return _fallback_load_docx(file_path)


def _load_pptx(file_path: str) -> list:
    try:
        documents = _run_with_timeout(lambda: UnstructuredPowerPointLoader(file_path).load())
        if any(d.page_content.strip() for d in documents):
            return documents
        raise ValueError("Unstructured loader returned empty content.")
    except Exception as e:
        print(f"⚠️ UnstructuredPowerPointLoader failed/timed out, using fallback parser: {e}")
        return _fallback_load_pptx(file_path)


def _load_excel(file_path: str) -> list:
    try:
        documents = _run_with_timeout(lambda: UnstructuredExcelLoader(file_path).load())
        if any(d.page_content.strip() for d in documents):
            return documents
        raise ValueError("Unstructured loader returned empty content.")
    except Exception as e:
        print(f"⚠️ UnstructuredExcelLoader failed/timed out, using fallback parser: {e}")
        return _fallback_load_excel(file_path)


def ingest_document(file_path: str, ext: str, user_id: int) -> int:
    """
    Parse, chunk, embed and store a document.
    Tries LangChain/Unstructured loaders first; falls back to lightweight
    glue-code parsers (python-docx/python-pptx/pandas) only if those fail.
    Returns the number of chunks stored.
    """
    documents = []
    filename = os.path.basename(file_path)

    if ext == ".pdf":
        try:
            loader = PyPDFLoader(file_path)
            documents = loader.load()
            if not any(doc.page_content.strip() for doc in documents):
                print(f"⚠️ PDF text empty, falling back to OCR: {filename}")
                return ingest_ocr_document(file_path, user_id)
        except Exception as e:
            print(f"⚠️ PDF parsing failed, falling back to OCR: {e}")
            return ingest_ocr_document(file_path, user_id)

    elif ext in [".txt", ".py", ".js", ".md", ".json"]:
        try:
            loader = TextLoader(file_path, encoding="utf-8")
            documents = loader.load()
        except UnicodeDecodeError:
            loader = TextLoader(file_path, encoding="latin-1")
            documents = loader.load()

    elif ext == ".docx":
        documents = _load_docx(file_path)
        if not documents:
            raise ValueError("No readable text found in this DOCX file.")

    elif ext == ".doc":
        # Legacy binary .doc has no reliable text loader — goes straight to OCR
        return ingest_ocr_document(file_path, user_id)

    elif ext == ".pptx":
        documents = _load_pptx(file_path)
        if not documents:
            raise ValueError("No readable text found in this PPTX file.")

    elif ext in [".xls", ".xlsx"]:
        documents = _load_excel(file_path)
        if not documents:
            raise ValueError("No readable data found in this spreadsheet.")

    elif ext in [".png", ".jpg", ".jpeg", ".bmp", ".tiff"]:
        return ingest_ocr_document(file_path, user_id)

    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if not documents:
        raise ValueError("No content could be extracted from this file.")

    for doc in documents:
        doc.metadata["source"] = "document"
        doc.metadata["file_type"] = ext
        doc.metadata["filename"] = filename

    splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50)
    chunks = splitter.split_documents(documents)

    if not chunks:
        raise ValueError("Document content was too short to process.")

    get_vector_store(user_id).add_documents(chunks)
    print(f"✅ Document ingested: {filename} → {len(chunks)} chunks")
    return len(chunks)